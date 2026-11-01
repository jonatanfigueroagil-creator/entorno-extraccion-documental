"""Fixtures comunes: PDF sintéticos creados al vuelo y localización opcional de los datos de práctica."""
from __future__ import annotations

import os
from pathlib import Path

import pytest
from pypdf import PdfWriter


def crear_pdf_con_texto(ruta: Path, texto: str) -> Path:
    """Escribe un PDF de una página con ``texto`` seleccionable (sin dependencias extra)."""
    flujo = f"BT /F1 14 Tf 72 720 Td ({texto}) Tj ET".encode("latin-1")
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(flujo)).encode() + b" >>\nstream\n" + flujo + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    salida = bytearray(b"%PDF-1.4\n")
    posiciones = []
    for numero, cuerpo in enumerate(objetos, start=1):
        posiciones.append(len(salida))
        salida += f"{numero} 0 obj\n".encode() + cuerpo + b"\nendobj\n"
    inicio_xref = len(salida)
    salida += f"xref\n0 {len(objetos) + 1}\n".encode() + b"0000000000 65535 f \n"
    for pos in posiciones:
        salida += f"{pos:010d} 00000 n \n".encode()
    salida += f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{inicio_xref}\n%%EOF\n".encode()
    ruta.write_bytes(bytes(salida))
    return ruta


def crear_pdf_sin_texto(ruta: Path, paginas: int = 1) -> Path:
    """Escribe un PDF de páginas en blanco (equivale a un escaneado sin capa de texto)."""
    escritor = PdfWriter()
    for _ in range(paginas):
        escritor.add_blank_page(width=612, height=792)
    with ruta.open("wb") as f:
        escritor.write(f)
    return ruta


@pytest.fixture
def lote_sintetico(tmp_path: Path) -> Path:
    """Carpeta con un PDF nativo, un PDF sin texto, un PDF dañado, un .eml y una imagen."""
    crear_pdf_con_texto(tmp_path / "A_nativo.pdf", "Estado de situacion financiera DEMO al 31 de diciembre")
    crear_pdf_sin_texto(tmp_path / "B_escaneado.pdf", paginas=2)
    (tmp_path / "C_danado.pdf").write_bytes(b"esto no es un pdf")
    (tmp_path / "D_reclamo.eml").write_text("Subject: Reclamo DEMO\n\nTexto", encoding="utf-8")
    (tmp_path / "E_foto.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    return tmp_path


@pytest.fixture(scope="session")
def datos_practica() -> Path:
    """Carpeta de documentos de práctica: variable SBS_DATOS_PRACTICA o la carpeta 02_Datos_Practica más cercana."""
    valor = os.environ.get("SBS_DATOS_PRACTICA")
    if valor and Path(valor).is_dir():
        return Path(valor)
    for carpeta in Path(__file__).resolve().parents:
        candidata = carpeta / "02_Datos_Practica"
        if candidata.is_dir():
            return candidata
    pytest.skip("No se encontraron los datos de práctica (define SBS_DATOS_PRACTICA)")


# ---------------------------------------------------------------- archivos sintéticos de S02

def crear_docx(ruta: Path, parrafos: list[str], tabla: list[list[str]] | None = None, combinar: bool = False) -> Path:
    """Escribe un .docx con párrafos y, opcionalmente, una tabla (con la primera fila combinada)."""
    from docx import Document

    documento = Document()
    for parrafo in parrafos:
        documento.add_paragraph(parrafo)
    if tabla:
        t = documento.add_table(rows=len(tabla), cols=len(tabla[0]))
        for i, fila in enumerate(tabla):
            for j, valor in enumerate(fila):
                t.cell(i, j).text = valor
        if combinar:
            t.cell(0, 0).merge(t.cell(0, len(tabla[0]) - 1))
    documento.save(str(ruta))
    return ruta


def crear_eml(ruta: Path, asunto: str, cuerpo: str, adjunto: tuple[str, bytes] | None = None) -> Path:
    """Escribe un correo .eml en UTF-8 (quoted-printable) con un adjunto opcional."""
    from email.message import EmailMessage

    mensaje = EmailMessage()
    mensaje["Subject"] = asunto
    mensaje["From"] = "ana@ejemplo.pe"
    mensaje["To"] = "reclamos@ejemplo.pe"
    mensaje.set_content(cuerpo, charset="utf-8", cte="quoted-printable")
    if adjunto:
        mensaje.add_attachment(adjunto[1], maintype="application", subtype="octet-stream", filename=adjunto[0])
    ruta.write_bytes(bytes(mensaje))
    return ruta


def crear_html(ruta: Path, cuerpo: str, codificacion: str = "utf-8", declarar: bool = True) -> Path:
    """Escribe un HTML con script y estilo; ``declarar=False`` omite el charset (caso latin-1)."""
    meta = f'<meta charset="{codificacion}">' if declarar else ""
    html = (f"<!DOCTYPE html><html><head>{meta}<title>t</title><style>p{{color:red}}</style>"
            f"<script>var x = 'no soy texto';</script></head><body><p>{cuerpo}</p></body></html>")
    ruta.write_bytes(html.encode(codificacion))
    return ruta


def crear_png(ruta: Path, lado: int = 32) -> Path:
    """Escribe un PNG de color liso (imagen sin texto)."""
    from PIL import Image

    Image.new("RGB", (lado, lado), "white").save(ruta, format="PNG")
    return ruta


def crear_pdf_cifrado(ruta: Path) -> Path:
    """Escribe un PDF protegido con contraseña (AES-256) usando PyMuPDF."""
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Documento reservado DEMO")
    doc.save(str(ruta), encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="clave", owner_pw="duena")
    doc.close()
    return ruta


@pytest.fixture
def lote_trampas(tmp_path: Path) -> Path:
    """Lote de 12 archivos con los 6 modos de fallo del caso: extensión falsa, DOCX corrupto, PDF cifrado,
    PDF truncado, archivo vacío y duplicado exacto; más un escaneado, un DOCX, un correo, un HTML y una imagen."""
    lote = tmp_path / "lote"
    lote.mkdir()
    nativo = crear_pdf_con_texto(lote / "A01_nativo.pdf", "Estado de situacion financiera DEMO al 31 de diciembre de 2026")
    crear_pdf_sin_texto(lote / "A02_escaneado.pdf", paginas=2)
    crear_png(lote / "A03_extension_falsa.pdf")
    (lote / "A04_docx_corrupto.docx").write_bytes(b"PK\x03\x04" + b"\x00" * 40)
    crear_pdf_cifrado(lote / "A05_cifrado.pdf")
    (lote / "A06_truncado.pdf").write_bytes(nativo.read_bytes()[:150])
    (lote / "A07_vacio.pdf").write_bytes(b"")
    (lote / "A08_copia.pdf").write_bytes(nativo.read_bytes())
    crear_docx(lote / "A09_contrato.docx", ["Contrato DEMO de servicios"], [["Cuota", "Monto"], ["1", "S/ 100,00"]])
    crear_eml(lote / "A10_reclamo.eml", "Reclamo DEMO", "Descuento indebido de la cuota.")
    crear_html(lote / "A11_tarifario.html", "Tasas y comisiones vigentes")
    crear_png(lote / "A12_voucher.png")
    return lote
